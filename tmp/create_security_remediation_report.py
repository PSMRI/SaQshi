from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('output/docs/SaQshi_Security_Remediation_Report.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)

def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), color); tc_pr.append(shd)

def cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(9)
    if color: r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; shade(c, '17365D'); cell_text(c, h, True, (255,255,255))
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            if n % 2: shade(cells[i], 'EAF1F8')
            cell_text(cells[i], val)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return t

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.add_run(text)
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(5)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet'); p.add_run(text); p.paragraph_format.space_after = Pt(2)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10)
for s in ['Title','Heading 1','Heading 2']:
    styles[s].font.name = 'Aptos'; styles[s].font.color.rgb = RGBColor(0,0,0)
styles['Title'].font.size = Pt(24); styles['Heading 1'].font.size = Pt(15); styles['Heading 2'].font.size = Pt(11)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.add_run('SaQshi Security Remediation Report')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Acunetix and Burp Findings'); r.bold=True; r.font.size=Pt(13)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Application: nhmsaqshi.piramalswasthya.org\nPrepared for security audit follow-up\nDate: 14 September 2026').font.size=Pt(10)

heading(doc, 'Executive Summary')
doc.add_paragraph('SaQshi has been updated to address the application and IIS configuration findings identified in the supplied Acunetix and Burp reports. The changes preserve existing authentication, CSRF, session, RBAC, assessment, dashboard and API behaviour. Technical deployment documentation is no longer exposed through the production web root; the built-in User Manual remains available.')
doc.add_paragraph('Several findings require deployment of the updated application root configuration. TLS protocol and cipher remediation additionally requires a Windows Server/IIS administrator change during an approved maintenance window. Those server-level changes are documented but have not been executed from the source repository.')

heading(doc, 'Implemented Remediations')
table(doc, ['Finding', 'Implemented remediation', 'Evidence'], [
('Password submitted using GET', 'Login form explicitly uses POST. JavaScript continues to prevent native submission and calls the existing JSON login API.', 'ui/pages/login/login.html'),
('Clickjacking and frameable responses', 'Global X-Frame-Options DENY and CSP frame-ancestors none are configured for IIS-served static content and API responses.', 'web.config; api/core/Security.php'),
('HSTS not enforced', 'HTTP requests permanently redirect to HTTPS. HSTS is added only to HTTPS responses through an IIS outbound rule.', 'web.config'),
('CSP missing on static responses', 'Global CSP covers self-hosted scripts, styles, fonts, images and API connections while retaining the inline sources currently required by the UI.', 'web.config'),
('Legacy XSS filter disabled', 'The obsolete X-XSS-Protection header is removed rather than enabled. CSP, output encoding and nosniff remain the protections.', 'web.config; api/core/Security.php'),
('Cacheable sensitive API content', 'API security headers use no-store, no-cache, must-revalidate, Pragma no-cache and Expires 0. Static assets remain cacheable.', 'api/core/Security.php'),
('HTML charset missing', 'IIS MIME mappings specify UTF-8 for static HTML and JSON responses.', 'web.config'),
('Technology/version disclosure', 'IIS removes X-Powered-By and server version where supported; PHP expose_php is disabled.', 'web.config; .user.ini'),
], [1.55, 3.65, 1.3])

heading(doc, 'Documentation and External Resource Controls')
table(doc, ['Area', 'Remediation'], [
('Public technical documentation', 'The production configuration hides the docs directory, GitBook reader and readme route. The built-in User Manual remains the user-facing documentation surface.'),
('Filesystem path disclosure', 'Public deployment examples no longer contain the reported Windows or Linux absolute paths.'),
('Credential-like examples', 'Database examples use explicit placeholder values such as <DB_USERNAME> and <DB_PASSWORD>.'),
('Third-party Swagger resources', 'Swagger UI and js-yaml assets are bundled locally, removing the external CDN/SRI finding.'),
('GitBook navigation', 'Public GitBook links were removed and replaced with the built-in User Manual.'),
], [2.1, 4.4])

heading(doc, 'Files Changed')
table(doc, ['File', 'Purpose'], [
('web.config', 'Global IIS security headers, HTTP to HTTPS redirect, HTTPS-only HSTS, CSP, UTF-8 mappings and production documentation restrictions.'),
('api/core/Security.php', 'API no-cache response controls and removal of obsolete X-XSS-Protection emission.'),
('ui/pages/login/login.html', 'Explicit POST form method and standard current-password autocomplete semantics.'),
('.user.ini', 'Disables PHP error display and PHP version exposure while retaining server logging.'),
('index.html', 'Replaces GitBook navigation with the built-in User Manual.'),
('docs/api/swagger-ui.html and docs/api/assets', 'Uses locally hosted Swagger support assets.'),
('docs/deployment/iis_tls_hardening.md', 'Windows/IIS SCHANNEL remediation procedure for TLS 1.0, TLS 1.1 and 3DES.'),
], [2.55, 3.95])

heading(doc, 'Server Administrator Actions Required')
doc.add_paragraph('The following actions are outside application-source control and must be completed on the production Windows Server/IIS host:')
bullet(doc, 'Deploy the updated application files, especially web.config, .user.ini and api/core/Security.php, then recycle the IIS application pool or site.')
bullet(doc, 'Disable TLS 1.0 and TLS 1.1 for both client and server SCHANNEL roles.')
bullet(doc, 'Disable 3DES and TLS_RSA_WITH_3DES_EDE_CBC_SHA; retain TLS 1.2 and enable TLS 1.3 where supported.')
bullet(doc, 'Confirm cipher-suite ordering follows the approved organisational Windows security baseline and favours AES-GCM or ChaCha20 suites.')
bullet(doc, 'Test the application after deployment before closing the audit findings.')

heading(doc, 'Findings Reviewed and Accepted')
table(doc, ['Finding', 'Disposition and justification'], [
('Credit card numbers disclosed', 'False positive. The reported numeric strings are decimal longitude/latitude portions in the Jharkhand GeoJSON boundary response, not payment-card data. Valid geographic coordinates were not modified.'),
('Email addresses disclosed', 'Authorized assessor and state-user APIs require email fields for their role-restricted administration workflows. Public contact email is intentional. No unnecessary public email exposure was identified in the remediated public routes.'),
('CSP unsafe-inline', 'Temporarily retained for existing inline UI scripts and styles. Removing it requires a controlled frontend refactor and full regression testing; unsafe-eval and wildcard sources are not used.'),
], [2.2, 4.3])

heading(doc, 'Verification Performed')
table(doc, ['Check', 'Result'], [
('web.config XML validation', 'Passed.'),
('PHP syntax validation for api/core/Security.php', 'Passed.'),
('Login form source check', 'Confirmed explicit POST method and current-password autocomplete.'),
('Legacy X-XSS-Protection emission', 'Confirmed absent from application code; IIS configuration explicitly removes inherited versions.'),
('Documentation source scan', 'Confirmed removal of scanner-identified absolute paths and credential-like defaults.'),
('Live production header check', 'Pre-deployment check showed the public server was still serving the prior IIS configuration. Deployment and re-scan remain required.'),
], [2.5, 4.0])

heading(doc, 'Security Team Reverification')
doc.add_paragraph('After deployment, verify the following endpoints: /, /ui/login.html, /ui/pages/login/login.html, /ui/dashboard.html, /api/auth/v1/csrf.php and /api/auth/v1/me.php.')
bullet(doc, 'HTTPS responses must include Strict-Transport-Security, X-Frame-Options DENY, X-Content-Type-Options nosniff, Referrer-Policy, Permissions-Policy and Content-Security-Policy.')
bullet(doc, 'Responses must not include X-XSS-Protection, X-Powered-By or detailed server/PHP version information where IIS supports removal.')
bullet(doc, 'HTTP requests must return a permanent redirect to the HTTPS equivalent.')
bullet(doc, 'Authenticated API responses must be no-store; CSS, JavaScript, icon fonts and public images may remain cacheable.')
bullet(doc, 'Re-run Burp and Acunetix after the IIS/TLS maintenance window and record the result with the deployed release version.')

footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('SaQshi Security Remediation Report | 14 September 2026').font.size = Pt(8)
doc.core_properties.title = 'SaQshi Security Remediation Report'
doc.core_properties.author = 'SaQshi Project Team'
doc.save(OUT)
print(OUT)
